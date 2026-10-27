# Shared python-docx helpers (맑은 고딕, 연구노트 스타일)
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

FONT = "맑은 고딕"
RED = RGBColor(0xC0, 0x00, 0x00)

def _font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia'):
        rf.set(qn(a), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color is not None: run.font.color.rgb = color

def rich(p, text, size=10, bold=False):
    """〔 〕 placeholders rendered red."""
    for part in re.split(r'(〔[^〕]*〕)', text):
        if not part: continue
        r = p.add_run(part)
        _font(r, size, bold, RED if part.startswith('〔') else None)
    return p

def new_doc(footer_text, landscape=False):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2); s.top_margin = Cm(2); s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.paragraph_format.space_after = Pt(3); st.paragraph_format.line_spacing = 1.15
    set_footer(s, footer_text)
    return d

def set_footer(section, text):
    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(fp.runs): r._element.getparent().remove(r._element)
    _font(fp.add_run(text + " · "), 8)
    r = fp.add_run(); _font(r, 8)
    for kind, val in (('begin', None), ('instr', 'PAGE'), ('end', None)):
        if kind == 'instr':
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = val
        else:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind)
        r._element.append(el)
    _font(fp.add_run(" 쪽"), 8)

def landscape_section(d):
    from docx.enum.section import WD_SECTION
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(1.5); s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.5)
    return s

def portrait_section(d):
    from docx.enum.section import WD_SECTION
    s = d.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2); s.top_margin = Cm(2); s.bottom_margin = Cm(1.8)
    return s

def para(d, text, size=10, bold=False, align=None, after=3, before=0, indent=None):
    p = d.add_paragraph(); rich(p, text, size, bold)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent: p.paragraph_format.left_indent = Cm(indent)
    return p

def h1(d, text):
    p = para(d, text, 13, True, before=10, after=5); p.paragraph_format.keep_with_next = True; return p

def h2(d, text):
    p = para(d, text, 11, True, before=6, after=3); p.paragraph_format.keep_with_next = True; return p

def bullet(d, text, size=10, level=0):
    p = para(d, "• " + text, size, indent=0.4 + 0.5 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    return p

def note(d, text):
    return para(d, "※ " + text, 8.5)

def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def _shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    tcPr.append(sh)

def table(d, rows, widths, header=True, size=9, label_col=False, row_h=None, center_cols=()):
    t = d.add_table(rows=len(rows), cols=len(widths))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, row in enumerate(rows):
        for j, w in enumerate(widths):
            c = t.cell(i, j); c.width = Cm(w)
            txt = row[j] if j < len(row) else ""
            c.paragraphs[0].text = ""
            lines = str(txt).split("\n")
            for k, line in enumerate(lines):
                p = c.paragraphs[0] if k == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.05
                hdr = (header and i == 0) or (label_col and j == 0)
                rich(p, line, size, hdr)
                if (header and i == 0) or j in center_cols: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if (header and i == 0) or (label_col and j == 0): _shade(c, 'EDEDED')
        if row_h:
            tr = t.rows[i]._tr; trPr = tr.get_or_add_trPr()
            h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(int(row_h * 567))); h.set(qn('w:hRule'), 'atLeast'); trPr.append(h)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 567)))
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def sign_block(d, writer="", checker="하현제"):
    table(d, [["작성자", "서명", "작성일", "점검자(연구책임자)", "서명", "점검일"],
              [writer, "", "", checker, "", ""]], [2.6, 2.4, 2.6, 3.4, 2.4, 2.6], size=9, row_h=0.9)

def manual_toc(d, entries):
    """수동 목차: Paragraph + 오른쪽 탭(점선 리더). entries = [(level, text, page)]"""
    for lvl, text, pg in entries:
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.paragraph_format.left_indent = Cm(0.6 * lvl); p.paragraph_format.space_after = Pt(2)
        _font(p.add_run(f"{text}\t{pg}"), 10 if lvl == 0 else 9.5, lvl == 0)
